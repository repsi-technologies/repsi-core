"""
Business reports computed from the gym's own records. Every report returns the same shape:

    {"type", "title", "filters", "columns": [{"key", "label", "kind"}], "rows": [...],
     "summary": {...}, "generated_at"}

so one table/export renderer works for all of them. Exports: CSV (built in), Excel (openpyxl) and
PDF (reportlab). Dates are the gym's local business days; money is in the gym's currency.
"""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from typing import Callable, Optional
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.attendance import Attendance
from app.models.booking import ClassBooking, ClassSession
from app.models.finance import Expense, Invoice, Payment, PaymentStatus, Refund
from app.models.member import Member, MemberStatus, Membership, MembershipPlan
from app.models.member_records import MembershipEvent
from app.models.trainer import GymClass, SessionStatus, Trainer, TrainerClient, TrainerClientStatus, TrainingSession
from app.models.user import Workspace

MAX_RANGE_DAYS = 731
GROUPINGS = ("day", "week", "month")


class ReportFilters:
    def __init__(self, start: date, end: date, group_by: str = "month", plan_id: Optional[str] = None,
                 branch_id: Optional[str] = None, trainer_id: Optional[str] = None):
        if end < start:
            raise HTTPException(status_code=422, detail="The end date must be on or after the start date.")
        if (end - start).days > MAX_RANGE_DAYS:
            raise HTTPException(status_code=422, detail="Reports cover at most two years at a time.")
        if group_by not in GROUPINGS:
            raise HTTPException(status_code=422, detail=f"group_by must be one of: {', '.join(GROUPINGS)}.")
        self.start, self.end, self.group_by = start, end, group_by
        self.plan_id, self.branch_id, self.trainer_id = plan_id, branch_id, trainer_id

    def as_dict(self) -> dict:
        return {k: (v.isoformat() if isinstance(v, date) else v) for k, v in vars(self).items() if v is not None}


def _tz(db: Session, workspace_id: str) -> ZoneInfo:
    ws = db.get(Workspace, workspace_id)
    try:
        return ZoneInfo(ws.timezone) if ws and ws.timezone else ZoneInfo("Asia/Kolkata")
    except Exception:
        return ZoneInfo("Asia/Kolkata")


def _bounds(tz: ZoneInfo, f: ReportFilters) -> tuple[datetime, datetime]:
    lo = datetime.combine(f.start, time.min, tzinfo=tz).astimezone(timezone.utc)
    hi = datetime.combine(f.end + timedelta(days=1), time.min, tzinfo=tz).astimezone(timezone.utc)
    return lo, hi


def _naive(dt: datetime) -> datetime:
    return dt.replace(tzinfo=None)


def _local(dt: datetime, tz: ZoneInfo) -> datetime:
    return (dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)).astimezone(tz)


def _bucket(d: date, group_by: str) -> str:
    if group_by == "day":
        return d.isoformat()
    if group_by == "week":
        monday = d - timedelta(days=d.weekday())
        return f"{monday.isoformat()}"
    return d.strftime("%Y-%m")


def _buckets(f: ReportFilters) -> list[str]:
    out, d = [], f.start
    while d <= f.end:
        b = _bucket(d, f.group_by)
        if not out or out[-1] != b:
            out.append(b)
        d += timedelta(days=1)
    return out


def _report(kind: str, title: str, f: ReportFilters, columns: list[tuple[str, str, str]], rows: list[dict],
            summary: dict) -> dict:
    return {"type": kind, "title": title, "filters": f.as_dict(),
            "columns": [{"key": k, "label": label, "kind": t} for k, label, t in columns],
            "rows": rows, "summary": summary, "generated_at": datetime.now(timezone.utc).isoformat()}


# ─── Revenue ───────────────────────────────────────────────────────────


def revenue(db: Session, ws: str, f: ReportFilters) -> dict:
    tz = _tz(db, ws)
    lo, hi = _bounds(tz, f)
    payments = db.scalars(select(Payment).where(
        Payment.workspace_id == ws, Payment.paid_at >= _naive(lo), Payment.paid_at < _naive(hi),
        Payment.status.in_((PaymentStatus.SUCCESS, PaymentStatus.REFUNDED)))).all()
    refunds = db.scalars(select(Refund).where(Refund.workspace_id == ws, Refund.processed_at >= lo, Refund.processed_at < hi)).all()
    expenses = db.scalars(select(Expense).where(Expense.workspace_id == ws, Expense.expense_date >= f.start,
                                                Expense.expense_date <= f.end)).all()
    rows = {b: {"period": b, "collected": 0.0, "refunds": 0.0, "net": 0.0, "expenses": 0.0, "profit": 0.0, "payments": 0}
            for b in _buckets(f)}
    by_method: dict[str, float] = defaultdict(float)
    for p in payments:
        b = _bucket(_local(p.paid_at, tz).date(), f.group_by)
        if b in rows:
            rows[b]["collected"] += p.amount
            rows[b]["payments"] += 1
            by_method[p.method.value if hasattr(p.method, "value") else str(p.method)] += p.amount
    for r in refunds:
        b = _bucket(_local(r.processed_at, tz).date(), f.group_by)
        if b in rows:
            rows[b]["refunds"] += r.amount
    for e in expenses:
        b = _bucket(e.expense_date, f.group_by)
        if b in rows:
            rows[b]["expenses"] += e.amount
    for r in rows.values():
        r["net"] = r["collected"] - r["refunds"]
        r["profit"] = r["net"] - r["expenses"]
        for k in ("collected", "refunds", "net", "expenses", "profit"):
            r[k] = round(r[k], 2)
    out = list(rows.values())
    total = lambda k: round(sum(r[k] for r in out), 2)  # noqa: E731
    return _report("revenue", "Revenue and collections", f, [
        ("period", "Period", "text"), ("payments", "Payments", "number"), ("collected", "Collected", "money"),
        ("refunds", "Refunds", "money"), ("net", "Net revenue", "money"), ("expenses", "Expenses", "money"),
        ("profit", "Profit", "money")], out, {
        "collected": total("collected"), "refunds": total("refunds"), "net": total("net"),
        "expenses": total("expenses"), "profit": total("profit"),
        "by_method": {k: round(v, 2) for k, v in sorted(by_method.items())},
        "average_payment": round(total("collected") / max(1, sum(r["payments"] for r in out)), 2)})


# ─── Memberships ───────────────────────────────────────────────────────


def memberships(db: Session, ws: str, f: ReportFilters) -> dict:
    q = select(Membership).where(Membership.workspace_id == ws)
    if f.plan_id:
        q = q.where(Membership.plan_id == f.plan_id)
    periods = db.scalars(q).all()
    renewed_ids = set(db.scalars(select(MembershipEvent.membership_id).where(
        MembershipEvent.workspace_id == ws, MembershipEvent.event_type == "renewed")).all())
    first_by_member: dict[str, date] = {}
    for m in periods:
        first_by_member[m.member_id] = min(first_by_member.get(m.member_id, m.start_date), m.start_date)
    rows = []
    for b in _buckets(f):
        if f.group_by == "month":
            y, mo = (int(x) for x in b.split("-"))
            b_start = date(y, mo, 1)
            b_end = (date(y + (mo == 12), mo % 12 + 1, 1) - timedelta(days=1))
        elif f.group_by == "week":
            b_start = date.fromisoformat(b)
            b_end = b_start + timedelta(days=6)
        else:
            b_start = b_end = date.fromisoformat(b)
        b_start, b_end = max(b_start, f.start), min(b_end, f.end)
        active_members = {m.member_id for m in periods if m.status != MemberStatus.CANCELLED
                          and m.start_date <= b_end and m.end_date >= b_start}
        new = {mid for mid, first in first_by_member.items() if b_start <= first <= b_end}
        renewals = sum(1 for m in periods if m.id in renewed_ids and b_start <= m.start_date <= b_end)
        expired = sum(1 for m in periods if b_start <= m.end_date <= b_end and m.status != MemberStatus.CANCELLED
                      and not any(n.member_id == m.member_id and n.start_date > m.end_date
                                  and n.start_date <= m.end_date + timedelta(days=30) for n in periods))
        cancelled = sum(1 for m in periods if m.status == MemberStatus.CANCELLED and b_start <= m.end_date <= b_end)
        rows.append({"period": b, "active": len(active_members), "new": len(new), "renewed": renewals,
                     "lapsed": expired, "cancelled": cancelled})
    by_plan: dict[str, int] = defaultdict(int)
    plans = {p.id: p.name for p in db.scalars(select(MembershipPlan).where(MembershipPlan.workspace_id == ws)).all()}
    today = date.today()
    for m in periods:
        if m.status == MemberStatus.ACTIVE and m.start_date <= today <= m.end_date:
            by_plan[plans.get(m.plan_id, "Unknown")] += 1
    ending = [m for m in periods if f.start <= m.end_date <= f.end and m.status != MemberStatus.CANCELLED]
    renewed_after = sum(1 for m in ending if any(n.member_id == m.member_id and n.start_date > m.start_date
                                                 and n.start_date <= m.end_date + timedelta(days=30) for n in periods))
    return _report("memberships", "Membership analytics", f, [
        ("period", "Period", "text"), ("active", "Active members", "number"), ("new", "New members", "number"),
        ("renewed", "Renewals", "number"), ("lapsed", "Lapsed (not renewed in 30 days)", "number"),
        ("cancelled", "Cancelled", "number")], rows, {
        "active_now_by_plan": dict(sorted(by_plan.items())),
        "ended_in_range": len(ending), "renewed_within_30_days": renewed_after,
        "renewal_rate_percent": round(100 * renewed_after / len(ending), 1) if ending else None,
        "new_members": sum(r["new"] for r in rows)})


# ─── Attendance ────────────────────────────────────────────────────────


def attendance(db: Session, ws: str, f: ReportFilters) -> dict:
    tz = _tz(db, ws)
    lo, hi = _bounds(tz, f)
    q = select(Attendance).where(Attendance.workspace_id == ws, Attendance.member_id.isnot(None),
                                 Attendance.check_in_time >= _naive(lo), Attendance.check_in_time < _naive(hi))
    if f.branch_id:
        q = q.where(Attendance.gym_location_id == f.branch_id)
    visits = db.scalars(q).all()
    rows = {b: {"period": b, "visits": 0, "unique_members": set(), "avg_minutes": []} for b in _buckets(f)}
    heat = [[0] * 24 for _ in range(7)]
    per_member: dict[str, int] = defaultdict(int)
    for v in visits:
        local = _local(v.check_in_time, tz)
        b = _bucket(local.date(), f.group_by)
        if b in rows:
            rows[b]["visits"] += 1
            rows[b]["unique_members"].add(v.member_id)
            if v.duration_minutes:
                rows[b]["avg_minutes"].append(v.duration_minutes)
        heat[local.weekday()][local.hour] += 1
        per_member[v.member_id] += 1
    out = []
    for r in rows.values():
        mins = r.pop("avg_minutes")
        out.append({**r, "unique_members": len(r["unique_members"]),
                    "avg_minutes": round(sum(mins) / len(mins)) if mins else None})
    peak = max(((d, h, heat[d][h]) for d in range(7) for h in range(24)), key=lambda x: x[2], default=None)
    names = {m.id: f"{m.first_name} {m.last_name or ''}".strip() for m in db.scalars(select(Member).where(
        Member.workspace_id == ws, Member.id.in_(list(per_member)))).all()} if per_member else {}
    top = sorted(per_member.items(), key=lambda x: -x[1])[:10]
    return _report("attendance", "Attendance analytics", f, [
        ("period", "Period", "text"), ("visits", "Visits", "number"), ("unique_members", "Members visiting", "number"),
        ("avg_minutes", "Average stay (min)", "number")], out, {
        "visits": len(visits), "unique_members": len(per_member),
        "visits_per_member": round(len(visits) / len(per_member), 1) if per_member else 0,
        "peak": {"weekday": ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")[peak[0]], "hour": peak[1],
                 "visits": peak[2]} if peak and peak[2] else None,
        "heatmap": heat,
        "top_members": [{"member_id": mid, "name": names.get(mid, ""), "visits": n} for mid, n in top]})


# ─── Dues ──────────────────────────────────────────────────────────────


def dues(db: Session, ws: str, f: ReportFilters) -> dict:
    from app.services.finance_ops import outstanding_dues

    items = outstanding_dues(db, ws, f.end)
    rows = [{k: d[k] for k in ("invoice_number", "member_name", "phone", "total", "paid", "balance", "due_date",
                                "days_overdue", "age_bucket")} for d in items]
    buckets: dict[str, float] = defaultdict(float)
    for d in items:
        buckets[d["age_bucket"]] += d["balance"]
    return _report("dues", "Outstanding dues", f, [
        ("invoice_number", "Invoice", "text"), ("member_name", "Member", "text"), ("phone", "Phone", "text"),
        ("total", "Total", "money"), ("paid", "Paid", "money"), ("balance", "Balance", "money"),
        ("due_date", "Due date", "date"), ("days_overdue", "Days overdue", "number"), ("age_bucket", "Age", "text")],
        rows, {"total_outstanding": round(sum(d["balance"] for d in items), 2),
               "by_age": {k: round(v, 2) for k, v in buckets.items()}, "invoices": len(items)})


# ─── Trainers ──────────────────────────────────────────────────────────


def trainers(db: Session, ws: str, f: ReportFilters) -> dict:
    tz = _tz(db, ws)
    lo, hi = _bounds(tz, f)
    q = select(Trainer).where(Trainer.workspace_id == ws, Trainer.is_active.is_(True))
    if f.trainer_id:
        q = q.where(Trainer.id == f.trainer_id)
    rows = []
    for t in db.scalars(q.order_by(Trainer.name)).all():
        clients = db.scalar(select(func.count(TrainerClient.id)).where(
            TrainerClient.workspace_id == ws, TrainerClient.trainer_id == t.id,
            TrainerClient.status == TrainerClientStatus.ACTIVE)) or 0
        appts = db.scalars(select(TrainingSession).where(
            TrainingSession.workspace_id == ws, TrainingSession.trainer_id == t.id,
            TrainingSession.scheduled_at >= lo, TrainingSession.scheduled_at < hi)).all()
        classes = db.scalars(select(ClassSession).where(
            ClassSession.workspace_id == ws, ClassSession.trainer_id == t.id, ClassSession.starts_at >= lo,
            ClassSession.starts_at < hi, ClassSession.status != "cancelled")).all()
        attendees = db.scalar(select(func.count(ClassBooking.id)).where(
            ClassBooking.session_id.in_([c.id for c in classes]), ClassBooking.status == "attended")) if classes else 0
        own_visits = db.scalar(select(func.count(Attendance.id)).where(
            Attendance.workspace_id == ws, Attendance.trainer_id == t.id, Attendance.check_in_time >= _naive(lo),
            Attendance.check_in_time < _naive(hi))) or 0
        completed = [a for a in appts if a.status == SessionStatus.COMPLETED]
        rows.append({
            "trainer": t.name, "active_clients": clients, "sessions_booked": len(appts),
            "sessions_completed": len(completed),
            "sessions_no_show": sum(1 for a in appts if a.status == SessionStatus.NO_SHOW),
            "classes_taught": len(classes), "class_attendees": attendees or 0, "days_present": own_visits,
            "pt_hours": round(sum((a.duration_minutes or 60) for a in completed) / 60, 1),
            "estimated_pt_earnings": round(sum((a.duration_minutes or 60) for a in completed) / 60 * (t.hourly_rate or 0), 2),
        })
    return _report("trainers", "Trainer performance", f, [
        ("trainer", "Trainer", "text"), ("active_clients", "Active clients", "number"),
        ("sessions_booked", "PT sessions", "number"), ("sessions_completed", "Completed", "number"),
        ("sessions_no_show", "No-shows", "number"), ("classes_taught", "Classes taught", "number"),
        ("class_attendees", "Class attendees", "number"), ("days_present", "Check-ins", "number"),
        ("pt_hours", "PT hours", "number"), ("estimated_pt_earnings", "PT earnings (hourly rate)", "money")],
        rows, {"trainers": len(rows), "pt_hours": round(sum(r["pt_hours"] for r in rows), 1)})


# ─── Classes ───────────────────────────────────────────────────────────


def classes(db: Session, ws: str, f: ReportFilters) -> dict:
    tz = _tz(db, ws)
    lo, hi = _bounds(tz, f)
    sessions = db.scalars(select(ClassSession).where(
        ClassSession.workspace_id == ws, ClassSession.starts_at >= lo, ClassSession.starts_at < hi)).all()
    by_class: dict[str, dict] = {}
    names = {c.id: c.name for c in db.scalars(select(GymClass).where(GymClass.workspace_id == ws)).all()}
    status_counts: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    if sessions:
        for sid, st, n in db.execute(select(ClassBooking.session_id, ClassBooking.status, func.count(ClassBooking.id)).where(
                ClassBooking.session_id.in_([s.id for s in sessions])).group_by(ClassBooking.session_id, ClassBooking.status)).all():
            status_counts[sid][st] = n
    for s in sessions:
        r = by_class.setdefault(s.class_id, {"class": names.get(s.class_id, "Class"), "sessions": 0, "cancelled": 0,
                                             "capacity": 0, "booked": 0, "attended": 0, "no_show": 0, "late_cancel": 0})
        if s.status == "cancelled":
            r["cancelled"] += 1
            continue
        r["sessions"] += 1
        r["capacity"] += s.capacity
        c = status_counts[s.id]
        r["attended"] += c.get("attended", 0)
        r["no_show"] += c.get("no_show", 0)
        r["booked"] += c.get("booked", 0) + c.get("attended", 0) + c.get("no_show", 0)
    late = db.execute(select(ClassSession.class_id, func.count(ClassBooking.id)).join(
        ClassSession, ClassSession.id == ClassBooking.session_id).where(
        ClassBooking.workspace_id == ws, ClassBooking.late_cancel.is_(True), ClassSession.starts_at >= lo,
        ClassSession.starts_at < hi).group_by(ClassSession.class_id)).all()
    for cid, n in late:
        if cid in by_class:
            by_class[cid]["late_cancel"] = n
    rows = []
    for r in by_class.values():
        r["fill_rate_percent"] = round(100 * r["booked"] / r["capacity"], 1) if r["capacity"] else None
        r["no_show_rate_percent"] = round(100 * r["no_show"] / r["booked"], 1) if r["booked"] else None
        rows.append(r)
    rows.sort(key=lambda r: r["class"])
    return _report("classes", "Class utilisation", f, [
        ("class", "Class", "text"), ("sessions", "Sessions", "number"), ("cancelled", "Cancelled", "number"),
        ("capacity", "Places", "number"), ("booked", "Booked", "number"), ("attended", "Attended", "number"),
        ("no_show", "No-shows", "number"), ("late_cancel", "Late cancels", "number"),
        ("fill_rate_percent", "Fill rate %", "number"), ("no_show_rate_percent", "No-show rate %", "number")],
        rows, {"sessions": sum(r["sessions"] for r in rows), "booked": sum(r["booked"] for r in rows)})


REPORTS: dict[str, Callable[[Session, str, ReportFilters], dict]] = {
    "revenue": revenue, "memberships": memberships, "attendance": attendance, "dues": dues,
    "trainers": trainers, "classes": classes,
}
CATALOG = [
    {"type": "revenue", "title": "Revenue and collections", "description": "Collections, refunds, expenses and profit over time, by payment method."},
    {"type": "memberships", "title": "Membership analytics", "description": "Active, new, renewing, lapsed and cancelled members; renewal rate; members by plan."},
    {"type": "attendance", "title": "Attendance analytics", "description": "Visits, unique members, average stay, peak hours heatmap and most active members."},
    {"type": "dues", "title": "Outstanding dues", "description": "Unpaid balances by member with ageing buckets."},
    {"type": "trainers", "title": "Trainer performance", "description": "Clients, PT sessions, classes taught, attendees and PT hours."},
    {"type": "classes", "title": "Class utilisation", "description": "Fill rate, attendance, no-shows and late cancellations per class."},
]


def run(db: Session, workspace_id: str, report_type: str, f: ReportFilters) -> dict:
    fn = REPORTS.get(report_type)
    if fn is None:
        raise HTTPException(status_code=404, detail="Unknown report type.")
    return fn(db, workspace_id, f)


# ─── Export ────────────────────────────────────────────────────────────


def _cell(v):
    if isinstance(v, (dict, list)):
        return str(v)
    return "" if v is None else v


def to_csv(report: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in report["columns"]])
    for r in report["rows"]:
        w.writerow([_safe_csv(_cell(r.get(c["key"]))) for c in report["columns"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _safe_csv(v):
    """Stops spreadsheet formula injection (cells starting with = + - @)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def to_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = report["type"][:31]
    ws.append([report["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{report['filters'].get('start', '')} to {report['filters'].get('end', '')}"])
    ws.append([])
    ws.append([c["label"] for c in report["columns"]])
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E25822")
    for r in report["rows"]:
        ws.append([_safe_csv(_cell(r.get(c["key"]))) for c in report["columns"]])
    for i, c in enumerate(report["columns"], start=1):
        if c["kind"] == "money":
            for row in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                row[0].number_format = "#,##0.00"
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = max(12, len(c["label"]) + 2)
    summary = wb.create_sheet("Summary")
    for k, v in report["summary"].items():
        if not isinstance(v, (dict, list)):
            summary.append([k.replace("_", " ").capitalize(), v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report: dict, gym_name: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=report["title"])
    esc = lambda t: str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    data = [[Paragraph(f"<b>{esc(c['label'])}</b>", styles["BodyText"]) for c in report["columns"]]]
    for r in report["rows"][:2000]:
        data.append([esc(_cell(r.get(c["key"]))) for c in report["columns"]])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FDF1EB")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E4E2DC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    summary = [f"{k.replace('_', ' ').capitalize()}: {v}" for k, v in report["summary"].items()
               if not isinstance(v, (dict, list))]
    story = [Paragraph(f"<b>{esc(gym_name)}</b> · {esc(report['title'])}", styles["Title"]),
             Paragraph(esc(f"{report['filters'].get('start', '')} to {report['filters'].get('end', '')}"), styles["Normal"]),
             Spacer(1, 4 * mm), table, Spacer(1, 4 * mm)]
    story += [Paragraph(esc(s), styles["Normal"]) for s in summary]
    doc.build(story)
    return buf.getvalue()
